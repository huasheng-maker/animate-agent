"""Deterministic file (pptx/docx/pdf/markdown/text) to DocumentIR parsers."""

from __future__ import annotations

import hashlib
import re
from pathlib import Path
from typing import Literal, cast

from markdown_it import MarkdownIt
from markdown_it.token import Token
from pydantic import JsonValue

from animate_agent.documents.models import DocumentBlock, DocumentIR, DocumentSource, Section

BlockKind = Literal[
    "paragraph",
    "code",
    "equation",
    "table",
    "image",
    "diagram",
    "list",
    "quote",
    "callout",
]

MARKDOWN_EXTENSIONS = frozenset({".md", ".markdown", ".txt"})
SUPPORTED_EXTENSIONS = frozenset({".pptx", ".docx", ".pdf"}) | MARKDOWN_EXTENSIONS

_FRONTMATTER_FENCE = ("---", "...")
_INLINE_TEXT_TOKENS = frozenset({"text", "code_inline"})

# Real documents routinely mark sections as 【一、基础知识】 or 一、基础知识 rather
# than with markdown `#`. Without these, the whole document collapses into one
# flat section and the section structure is lost before the LLM ever sees it.
_BRACKET_HEADING_RE = re.compile(r"^【\s*([^】]+?)\s*】$")
_NUMBERED_HEADING_RE = re.compile(
    r"^(?:第\s*[一二三四五六七八九十百零〇\d]+\s*[章节節讲部分篇]"
    r"|[一二三四五六七八九十]{1,3}\s*[、.．]"
    r"|[（(]\s*[一二三四五六七八九十]{1,3}\s*[）)])"
    r"\s*\S.*$"
)
_PLAIN_HEADING_MAX_LENGTH = 40
_TITLE_MAX_LENGTH = 60

_HUGO_COMMENT_RE = re.compile(
    r"\{\{[<%]\s*comment\s*[>%]\}\}.*?\{\{[<%]\s*/comment\s*[>%]\}\}",
    re.DOTALL | re.IGNORECASE,
)
_HUGO_TOOLTIP_RE = re.compile(
    r"\{\{[<%]\s*glossary_tooltip\s+(?P<attrs>.*?)\s*[>%]\}\}",
    re.DOTALL | re.IGNORECASE,
)
_HUGO_HEADING_RE = re.compile(
    r"\{\{%\s*heading\s+[\"'](?P<slug>[^\"']+)[\"']\s*%\}\}",
    re.IGNORECASE,
)
_HUGO_SHORTCODE_RE = re.compile(r"\{\{[<%].*?[>%]\}\}", re.DOTALL)
_HUGO_ATTR_RE = re.compile(r"(?P<name>[a-zA-Z_][\w-]*)=[\"'](?P<value>.*?)[\"']")


def _clean_text(value: str) -> str:
    return re.sub(r"\s+", " ", value).strip()


def _strip_hugo_shortcodes(value: str) -> str:
    """Turn common Kubernetes/Hugo authoring syntax into readable Markdown.

    Shortcodes are publishing instructions rather than document content. Keep
    visible tooltip text, retain note bodies, and drop author-only comments so
    neither DocumentIR nor the model sees template syntax as lesson material.
    """

    value = _HUGO_COMMENT_RE.sub("", value)

    def tooltip_text(match: re.Match[str]) -> str:
        attrs = {
            item.group("name"): item.group("value")
            for item in _HUGO_ATTR_RE.finditer(match.group("attrs"))
        }
        return attrs.get("text") or attrs.get("term_id", "")

    def heading_text(match: re.Match[str]) -> str:
        words = match.group("slug").replace("-", " ").strip()
        return words.capitalize()

    value = _HUGO_TOOLTIP_RE.sub(tooltip_text, value)
    value = _HUGO_HEADING_RE.sub(heading_text, value)
    return _HUGO_SHORTCODE_RE.sub("", value)


def _document_id(path: Path) -> str:
    return hashlib.sha256(str(path.resolve()).encode("utf-8")).hexdigest()[:16]


def _is_heading_style(style_name: str) -> bool:
    lowered = style_name.lower()
    return lowered.startswith("heading") or "标题" in style_name or lowered == "title"


def _paragraph_block(section_id: str, index: int, text: str) -> DocumentBlock | None:
    cleaned = _clean_text(text)
    if not cleaned:
        return None
    return DocumentBlock(id=f"{section_id}-block-{index}", type="paragraph", text=cleaned)


def _resolve_title(sections: list[Section]) -> str:
    for section in sections:
        if section.title:
            return section.title
    for section in sections:
        for block in section.blocks:
            if block.text:
                return block.text
    return "Untitled document"


def _finalize(path: Path, sections: list[Section], *, title: str | None = None) -> DocumentIR:
    return DocumentIR(
        document_id=_document_id(path),
        title=title or _resolve_title(sections),
        source=DocumentSource(type="file"),
        sections=sections,
    )


def _plain_heading(text: str) -> str | None:
    """Recognize a heading written in a common Chinese plain-text convention."""
    if len(text) > _PLAIN_HEADING_MAX_LENGTH:
        return None
    bracketed = _BRACKET_HEADING_RE.match(text)
    if bracketed:
        return bracketed.group(1).strip() or None
    if _NUMBERED_HEADING_RE.match(text):
        return text
    return None


def _promote_leading_title(sections: list[Section]) -> str | None:
    """Promote a leading title line to the document title and drop it from content.

    Plain-text documents often open with the title as a bare first line, before
    any heading. That line is metadata rather than a teachable block — and the
    Knowledge Agent prompt already carries the title separately, so leaving it in
    would only add a block that the source_refs coverage check then demands be
    cited. Only promotes when real structure follows the preamble.
    """
    if len(sections) < 2 or sections[0].id != "section-overview":
        return None
    blocks = sections[0].blocks
    if not blocks or blocks[0].type != "paragraph" or len(blocks[0].text) > _TITLE_MAX_LENGTH:
        return None
    title = blocks.pop(0).text
    if not blocks:
        sections.pop(0)
    return title


def parse_pptx(path: str | Path) -> DocumentIR:
    """Convert a PPTX presentation into DocumentIR (one Section per slide)."""
    from pptx import Presentation

    p = Path(path)
    presentation = Presentation(str(p))
    sections: list[Section] = []
    for slide_index, slide in enumerate(presentation.slides, start=1):
        lines: list[str] = []
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            for paragraph in shape.text_frame.paragraphs:
                text = "".join(run.text for run in paragraph.runs)
                cleaned = _clean_text(text)
                if cleaned:
                    lines.append(cleaned)
        if not lines:
            continue
        section = Section(id=f"section-{slide_index}", title=lines[0], level=1)
        for index, text in enumerate(lines[1:], start=1):
            block = _paragraph_block(section.id, index, text)
            if block is not None:
                section.blocks.append(block)
        sections.append(section)
    return _finalize(p, sections)


def parse_docx(path: str | Path) -> DocumentIR:
    """Convert a DOCX document into DocumentIR, grouping paragraphs by heading styles."""
    from docx import Document

    p = Path(path)
    document = Document(str(p))
    sections: list[Section] = []
    current: Section | None = None
    section_counter = 0

    for paragraph in document.paragraphs:
        style_name = paragraph.style.name if paragraph.style is not None else ""
        text = _clean_text(paragraph.text)
        if not text:
            continue
        if _is_heading_style(style_name):
            section_counter += 1
            current = Section(id=f"section-{section_counter}", title=text, level=1)
            sections.append(current)
            continue
        if current is None:
            section_counter += 1
            current = Section(id=f"section-{section_counter}", title="", level=1)
            sections.append(current)
        block = _paragraph_block(current.id, len(current.blocks) + 1, text)
        if block is not None:
            current.blocks.append(block)

    return _finalize(p, sections)


def parse_pdf(path: str | Path) -> DocumentIR:
    """Convert a PDF into DocumentIR (one Section per page)."""
    from pypdf import PdfReader

    p = Path(path)
    reader = PdfReader(str(p))
    sections: list[Section] = []
    for page_index, page in enumerate(reader.pages, start=1):
        raw = page.extract_text() or ""
        lines = [cleaned for line in raw.splitlines() if (cleaned := _clean_text(line))]
        if not lines:
            continue
        section = Section(id=f"section-{page_index}", title=lines[0], level=1)
        for index, text in enumerate(lines[1:], start=1):
            block = _paragraph_block(section.id, index, text)
            if block is not None:
                section.blocks.append(block)
        sections.append(section)
    return _finalize(p, sections)


def _strip_frontmatter(text: str) -> str:
    """Drop a leading YAML frontmatter block so its keys don't leak into the content."""
    lines = text.splitlines()
    if not lines or lines[0].strip() != "---":
        return text
    for index in range(1, len(lines)):
        if lines[index].strip() in _FRONTMATTER_FENCE:
            return "\n".join(lines[index + 1 :])
    return text


def _read_text(path: Path) -> str:
    try:
        return path.read_text(encoding="utf-8")
    except UnicodeDecodeError as exc:
        raise ValueError(f"文件不是 UTF-8 文本，无法解析: {path.name}") from exc


def _inline_text(token: Token) -> str:
    """Flatten an inline token's children to plain text, dropping markdown markers."""
    parts: list[str] = []
    for child in token.children or []:
        if child.type in _INLINE_TEXT_TOKENS:
            parts.append(child.content)
        elif child.type in {"softbreak", "hardbreak"}:
            parts.append(" ")
        elif child.type == "image":
            parts.append(child.content)  # alt text
    return _clean_text("".join(parts))


def parse_markdown(path: str | Path) -> DocumentIR:
    """Convert Markdown / plain text into DocumentIR.

    Used for both `.md` and `.txt`: markdown gives the structure, and plain text
    simply arrives as paragraphs under a single overview section. `.txt` authored
    with markdown-style headings therefore parses the same way.
    """
    p = Path(path)
    source = _strip_hugo_shortcodes(_strip_frontmatter(_read_text(p)))
    tokens = MarkdownIt("commonmark").enable("table").parse(source)
    sections: list[Section] = []
    current: Section | None = None
    section_counter = 0

    def target_section() -> Section:
        nonlocal current
        if current is None:
            current = Section(id="section-overview", title="", level=1)
            sections.append(current)
        return current

    def start_section(title: str, level: int) -> None:
        nonlocal current, section_counter
        # Count sections rather than reading len(sections): the synthetic
        # section-overview sits in that list too, and letting it consume a
        # number leaves a gap (section-2 onwards, with no section-1).
        section_counter += 1
        current = Section(id=f"section-{section_counter}", title=title, level=level)
        sections.append(current)

    def add_block(
        kind: BlockKind,
        text: str,
        *,
        language: str | None = None,
        source_ref: str | None = None,
        metadata: dict[str, JsonValue] | None = None,
    ) -> None:
        section = target_section()
        section.blocks.append(
            DocumentBlock(
                id=f"{section.id}-block-{len(section.blocks) + 1}",
                type=kind,
                text=text,
                language=language,
                source_ref=source_ref,
                metadata=metadata or {},
            )
        )

    index = 0
    while index < len(tokens):
        token = tokens[index]

        if token.type == "heading_open":
            title = _inline_text(tokens[index + 1])
            if title:
                start_section(title, int(token.tag[1]))
            index += 1
            continue

        if token.type == "fence":
            code = token.content.rstrip("\n")
            if code.strip():
                language = token.info.strip() or None
                kind: BlockKind = "equation" if language in {"math", "latex", "tex"} else "code"
                add_block(kind, code, language=language)
            index += 1
            continue

        if token.type == "blockquote_open":
            parts: list[str] = []
            index += 1
            while index < len(tokens) and tokens[index].type != "blockquote_close":
                if tokens[index].type == "inline":
                    text = _inline_text(tokens[index])
                    if text:
                        parts.append(text)
                index += 1
            text = " ".join(parts)
            if text:
                callout = re.match(r"^\[!(?P<kind>[A-Za-z]+)\]\s*(?P<body>.*)$", text)
                if callout:
                    add_block(
                        "callout",
                        callout.group("body") or callout.group("kind"),
                        metadata={"kind": callout.group("kind").lower()},
                    )
                else:
                    add_block("quote", text)
            index += 1
            continue

        if token.type == "table_open":
            rows: list[list[str]] = []
            row: list[str] | None = None
            index += 1
            while index < len(tokens) and tokens[index].type != "table_close":
                if tokens[index].type == "tr_open":
                    row = []
                elif tokens[index].type == "inline" and row is not None:
                    row.append(_inline_text(tokens[index]))
                elif tokens[index].type == "tr_close" and row:
                    rows.append(row)
                    row = None
                index += 1
            if rows:
                add_block(
                    "table",
                    "\n".join(" | ".join(cells) for cells in rows),
                    metadata={"rows": cast(JsonValue, rows)},
                )
            index += 1
            continue

        if token.type == "paragraph_open":
            inline_index = index + 1
            children = tokens[inline_index].children or []
            if len(children) == 1 and children[0].type == "image":
                # A paragraph that is nothing but an image is an image block.
                image = children[0]
                src = image.attrGet("src")
                add_block(
                    "image",
                    _clean_text(image.content),
                    source_ref=str(src) if src is not None else None,
                )
            else:
                text = _inline_text(tokens[inline_index])
                heading = _plain_heading(text) if text else None
                if heading is not None:
                    start_section(heading, 1)
                elif text:
                    add_block("paragraph", text)
            index += 1
            continue

        if token.type in {"bullet_list_open", "ordered_list_open"}:
            closing = token.type.replace("_open", "_close")
            index += 1
            while index < len(tokens) and tokens[index].type != closing:
                if tokens[index].type == "inline":
                    item = _inline_text(tokens[index])
                    if item:
                        # One block per item: a single blob for the whole list
                        # leaves source_refs unable to name the step a scene is
                        # actually about.
                        add_block("list", item)
                index += 1
            index += 1
            continue

        index += 1

    if not sections:
        raise ValueError(f"文件中没有可解析的文本内容: {p.name}")
    return _finalize(p, sections, title=_promote_leading_title(sections))


def parse_file(path: str | Path) -> DocumentIR:
    """Dispatch on file extension and return a DocumentIR."""
    p = Path(path)
    ext = p.suffix.lower()
    if ext == ".pptx":
        return parse_pptx(p)
    if ext == ".docx":
        return parse_docx(p)
    if ext == ".pdf":
        return parse_pdf(p)
    if ext in MARKDOWN_EXTENSIONS:
        return parse_markdown(p)
    supported = ", ".join(sorted(SUPPORTED_EXTENSIONS))
    raise ValueError(f"不支持的文件格式: {ext}。支持: {supported}")
